import shutil
import tempfile
from pathlib import Path

from src.ocr.application.count_keywords import CountKeywordsUseCase
from src.ocr.infrastructure.postprocessing.keyword_counter import KeywordCounter


SAMPLE_ANNUAL_REPORT_TEXT = """
BÁO CÁO THƯỜNG NIÊN NĂM 2025 - TẬP ĐOÀN CÔNG NGHỆ VÀ TÀI CHÍNH AN PHÁT

Trong năm tài chính 2025, Tập đoàn An Phát ghi nhận bước chuyển mình lịch sử khi đặt trọng tâm phát triển vào các trụ cột công nghệ cốt lõi. Ban điều hành xác định việc thúc đẩy trí tuệ nhân tạo là đòn bẩy chiến lược để tái định hình toàn bộ trải nghiệm khách hàng và tối ưu hóa chi phí vận hành. Hệ thống tự động hóa chăm sóc khách hàng dựa trên nền tảng trí tuệ nhân tạo chính thức đi vào hoạt động, giúp giảm 45% thời gian phản hồi giao dịch. 

Tuy nhiên, trong các tài liệu kiểm toán ban đầu, do quá trình quét scan ảnh chất lượng thấp, hệ thống OCR từng ghi nhận dạng biến thể tri tuệ nhan tạo và cả cụm từ trí tue nhân tao ở một số hợp đồng chi nhánh tỉnh. Với giải pháp mới, toàn bộ các biến thể này đều được khoanh vùng chuẩn xác.

Bên cạnh đó, chiến lược chuyển đổi số toàn diện tiếp tục được triển khai sâu rộng trên 63 tỉnh thành. Việc nâng cấp hạ tầng sang điện toán đám mây đã nâng cao tính sẵn sàng của hệ thống lên mức 99,99%. Một số văn bản lưu trữ cũ bị mờ nét thành chuyen doi so hay điên toan đám mây, nhưng năng lực phát hiện ngữ cảnh vẫn nhận diện trọn vẹn giá trị công nghệ mà tập đoàn đã đầu tư.

Đặc biệt, năng lực khai thác dữ liệu lớn trở thành tài sản vô giá của An Phát. Cơ sở hạ tầng dữ liệu lớn xử lý hơn 100 triệu giao dịch mỗi ngày. Mặc dù các bản scan scan-to-pdf xuất hiện lỗi chính tả OCR như du lieu lon và khai thac du lieu, trung tâm phân tích vẫn đảm bảo tính thống nhất của thông tin.

Về mảng học máy và học sâu, chúng tôi tự hào sở hữu đội ngũ kỹ sư hàng đầu. Các thuật toán học máy giúp phát hiện gian lận thanh toán theo thời gian thực. Một số biên bản đánh giá nội bộ bị lỗi gõ học may và học sầu cũng được liên kết đúng với các đề tài nghiên cứu tương ứng.

Cần lưu ý rằng, một số khái niệm có chữ tương tự nhưng không thuộc danh mục công nghệ mục tiêu như trí tuệ của nhân dân, dữ liệu không đầy đủ, hay đám mây trên trời trong báo cáo môi trường đã được hệ thống phân tách rạch ròi, tuyệt đối không bị đếm nhầm.

Tổng kết lại, công tác đầu tư công nghệ năm 2025 đã mang lại tỷ suất sinh lời vượt trội, tạo tiền đề vững chắc cho giai đoạn tăng trưởng bứt phá tiếp theo.
"""


def test_e2e_keyword_counter_with_complex_traps():
    # Target keywords under test
    target_keywords = [
        "trí tuệ nhân tạo",
        "chuyển đổi số",
        "điện toán đám mây",
        "dữ liệu lớn",
        "khai thác dữ liệu",
        "học máy",
        "học sâu",
    ]

    with tempfile.TemporaryDirectory() as tmpdir:
        base_dir = Path(tmpdir)
        doc_dir = base_dir / "AnPhat_BaoCaoThuongNien_2025"
        doc_dir.mkdir(parents=True, exist_ok=True)

        # 1. Write mock result.md
        result_md_path = doc_dir / "result.md"
        result_md_path.write_text(SAMPLE_ANNUAL_REPORT_TEXT, encoding="utf-8")

        # 2. Write custom keywords.txt
        kw_file = base_dir / "keywords_test.txt"
        kw_file.write_text("\n".join(target_keywords), encoding="utf-8")

        # ---------------------------------------------------------------------
        # Case A: Baseline thuần thuật toán (Exact Only: enable_llm=False)
        # ---------------------------------------------------------------------
        report_exact = CountKeywordsUseCase.execute(
            doc_dir=doc_dir,
            document_id="AnPhat_BaoCaoThuongNien_2025",
            keywords_file=str(kw_file),
            enable_llm=False,
        )

        assert report_exact is not None
        assert report_exact.llm_enabled is False

        # Exact counts expected in text:
        # "trí tuệ nhân tạo": xuất hiện 2 lần ("thúc đẩy trí tuệ nhân tạo", "nền tảng trí tuệ nhân tạo")
        # "chuyển đổi số": xuất hiện 1 lần
        # "điện toán đám mây": xuất hiện 1 lần
        # "dữ liệu lớn": xuất hiện 2 lần ("năng lực khai thác dữ liệu lớn", "hạ tầng dữ liệu lớn")
        # "khai thác dữ liệu": xuất hiện 1 lần (nằm trong "khai thác dữ liệu lớn")
        # "học máy": xuất hiện 2 lần ("mảng học máy", "thuật toán học máy")
        # "học sâu": xuất hiện 1 lần ("học sâu")
        exact_counts_map = {k.keyword: k.exact_count for k in report_exact.keywords}
        assert exact_counts_map["trí tuệ nhân tạo"] == 2
        assert exact_counts_map["chuyển đổi số"] == 1
        assert exact_counts_map["điện toán đám mây"] == 1
        assert exact_counts_map["dữ liệu lớn"] == 2
        assert exact_counts_map["học máy"] == 2
        assert exact_counts_map["học sâu"] == 1

        # ---------------------------------------------------------------------
        # Case B: Kiểm tra xuất file bảng tính Excel (count_keywords.xlsx)
        # ---------------------------------------------------------------------
        count_xlsx_path = doc_dir / "count_keywords.xlsx"
        assert count_xlsx_path.exists(), "File count_keywords.xlsx phải được tạo"

        import openpyxl
        wb = openpyxl.load_workbook(count_xlsx_path)
        ws = wb.active
        assert ws.title == "Thống kê từ khóa"
        assert ws["A1"].value == "BÁO CÁO TẦN SUẤT TỪ KHÓA"
        assert ws["A4"].value == "STT"
        assert ws["B4"].value == "Từ khóa / Cụm từ"
        assert ws["C4"].value == "Số lần xuất hiện"

        # Đọc dữ liệu từ file Excel
        excel_counts = {}
        for row in range(5, 5 + len(target_keywords)):
            kw = ws.cell(row=row, column=2).value
            cnt = ws.cell(row=row, column=3).value
            excel_counts[kw] = cnt

        assert excel_counts["trí tuệ nhân tạo"] == 2
        assert excel_counts["chuyển đổi số"] == 1
        assert excel_counts["điện toán đám mây"] == 1
        assert excel_counts["dữ liệu lớn"] == 2
        assert excel_counts["học máy"] == 2
        assert excel_counts["học sâu"] == 1

