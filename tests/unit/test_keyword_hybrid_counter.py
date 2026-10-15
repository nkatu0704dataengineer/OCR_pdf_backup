import tempfile
from pathlib import Path
import openpyxl

from src.ocr.domain.models.keyword import KeywordReport, KeywordCount
from src.ocr.infrastructure.postprocessing.keyword_counter import KeywordCounter
from src.ocr.infrastructure.export.keyword_exporter import KeywordExporter


def test_keyword_counter_tier1_pure_exact():
    text = (
        "Báo cáo triển khai trí tuệ nhân tạo kết hợp công nghệ mới và trí tuệ nhân tạo chuẩn. "
        "Đồng thời phát triển dữ liệu lớn và điện toán đám mây."
    )
    keywords = ["trí tuệ nhân tạo", "dữ liệu lớn", "học sâu"]

    report: KeywordReport = KeywordCounter.count_keywords(
        text=text,
        keywords=keywords,
        document_id="test_doc",
    )

    assert report.total_keywords_searched == 3
    assert report.total_exact_matches == 3  # 2 for 'trí tuệ nhân tạo', 1 for 'dữ liệu lớn'
    assert report.total_fuzzy_matches == 0
    assert report.total_matches == 3
    assert report.llm_enabled is False

    # Check breakdown for 'trí tuệ nhân tạo'
    ttnt = next(k for k in report.keywords if k.keyword == "trí tuệ nhân tạo")
    assert ttnt.exact_count == 2
    assert ttnt.count == 2

    # Check breakdown for 'dữ liệu lớn'
    dll = next(k for k in report.keywords if k.keyword == "dữ liệu lớn")
    assert dll.exact_count == 1
    assert dll.count == 1

    # Check breakdown for 'học sâu'
    hs = next(k for k in report.keywords if k.keyword == "học sâu")
    assert hs.exact_count == 0
    assert hs.count == 0


def test_keyword_exporter_excel_format():
    report = KeywordReport(
        document_id="demo_report",
        source_file="result.md",
        total_keywords_searched=2,
        total_exact_matches=5,
        total_fuzzy_matches=0,
        total_matches=5,
        llm_enabled=False,
        keywords=[
            KeywordCount(
                keyword="trí tuệ nhân tạo",
                exact_count=3,
                count=3,
            ),
            KeywordCount(
                keyword="chuyển đổi số",
                exact_count=2,
                count=2,
            )
        ]
    )

    with tempfile.TemporaryDirectory() as tmpdir:
        out_dir = Path(tmpdir)
        # Create a dummy count_keywords.md to test cleanup
        dummy_md = out_dir / "count_keywords.md"
        dummy_md.write_text("old markdown", encoding="utf-8")

        exported_path = KeywordExporter.export_excel(report, out_dir)
        assert exported_path.exists()
        assert exported_path.suffix == ".xlsx"
        # Old markdown file must be cleaned up
        assert not dummy_md.exists()

        # Verify Excel workbook contents
        wb = openpyxl.load_workbook(exported_path)
        ws = wb.active
        assert ws.title == "Thống kê từ khóa"
        assert ws["A1"].value == "BÁO CÁO TẦN SUẤT TỪ KHÓA"
        assert "demo_report" in ws["A2"].value
        assert ws["A4"].value == "STT"
        assert ws["B4"].value == "Từ khóa / Cụm từ"
        assert ws["C4"].value == "Số lần xuất hiện"
        assert ws["B5"].value == "trí tuệ nhân tạo"
        assert ws["C5"].value == 3
        assert ws["B6"].value == "chuyển đổi số"
        assert ws["C6"].value == 2
        assert ws["B7"].value == "Tổng cộng"
