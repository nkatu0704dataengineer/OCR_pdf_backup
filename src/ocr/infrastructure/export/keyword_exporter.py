import os
from pathlib import Path
from src.ocr.domain.models.keyword import KeywordReport

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False


class KeywordExporter:
    """
    Exports a KeywordReport as an Excel spreadsheet written to ``count_keywords.xlsx``
    in the same directory as ``result.md``.
    """

    @classmethod
    def export_excel(cls, report: KeywordReport, output_dir: Path) -> Path:
        """
        Writes ``count_keywords.xlsx`` inside *output_dir* and returns its path.
        Also cleans up any stale ``count_keywords.md`` in the same directory.
        """
        output_dir = Path(output_dir)
        output_dir.mkdir(parents=True, exist_ok=True)
        xlsx_path = output_dir / "count_keywords.xlsx"

        if not HAS_OPENPYXL:
            raise ImportError(
                "Thư viện 'openpyxl' chưa được cài đặt. Vui lòng chạy 'pip install openpyxl' để xuất file Excel."
            )

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Thống kê từ khóa"

        # 1. Report Title & Metadata
        ws["A1"] = "BÁO CÁO TẦN SUẤT TỪ KHÓA"
        ws["A1"].font = Font(name="Calibri", size=14, bold=True, color="1F4E79")
        ws.row_dimensions[1].height = 24

        meta_info = f"Tài liệu: {report.document_id}  |  Tổng số từ khóa tìm kiếm: {report.total_keywords_searched}  |  Tổng số lần xuất hiện: {report.total_matches}"
        ws["A2"] = meta_info
        ws["A2"].font = Font(name="Calibri", size=10, italic=True, color="595959")
        ws.row_dimensions[2].height = 18

        # 2. Table Header
        headers = ["STT", "Từ khóa / Cụm từ", "Số lần xuất hiện"]
        for col_idx, h_text in enumerate(headers, start=1):
            cell = ws.cell(row=4, column=col_idx, value=h_text)
            cell.fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
            cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        ws.row_dimensions[4].height = 25

        # Border styles
        cell_border = Border(
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9"),
            top=Side(style="thin", color="D9D9D9"),
            bottom=Side(style="thin", color="D9D9D9")
        )

        # 3. Data Rows
        start_row = 5
        for idx, kw in enumerate(report.keywords, start=1):
            curr_row = start_row + idx - 1
            ws.row_dimensions[curr_row].height = 20

            c_stt = ws.cell(row=curr_row, column=1, value=idx)
            c_kw = ws.cell(row=curr_row, column=2, value=kw.keyword)
            c_cnt = ws.cell(row=curr_row, column=3, value=kw.count)

            c_stt.alignment = Alignment(horizontal="center", vertical="center")
            c_kw.alignment = Alignment(horizontal="left", vertical="center")
            c_cnt.alignment = Alignment(horizontal="right", vertical="center")
            c_cnt.number_format = "#,##0"

            for c in (c_stt, c_kw, c_cnt):
                c.font = Font(name="Calibri", size=11)
                c.border = cell_border

        # 4. Total Summary Row
        tot_row = start_row + len(report.keywords)
        ws.row_dimensions[tot_row].height = 22

        ws.cell(row=tot_row, column=1, value="")
        c_label = ws.cell(row=tot_row, column=2, value="Tổng cộng")
        c_label.font = Font(name="Calibri", size=11, bold=True, color="1F4E79")
        c_label.alignment = Alignment(horizontal="right", vertical="center")

        formula_sum = f"=SUM(C{start_row}:C{tot_row - 1})" if len(report.keywords) > 0 else 0
        c_val = ws.cell(row=tot_row, column=3, value=formula_sum)
        c_val.font = Font(name="Calibri", size=11, bold=True, color="1F4E79")
        c_val.alignment = Alignment(horizontal="right", vertical="center")
        c_val.number_format = "#,##0"

        tot_border = Border(
            top=Side(style="thin", color="1F4E79"),
            bottom=Side(style="double", color="1F4E79"),
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9")
        )
        tot_fill = PatternFill(start_color="EDF2F8", end_color="EDF2F8", fill_type="solid")
        for col_idx in (1, 2, 3):
            cell = ws.cell(row=tot_row, column=col_idx)
            cell.border = tot_border
            cell.fill = tot_fill

        # 5. Auto-fit column widths
        max_kw_len = max((len(kw.keyword) for kw in report.keywords), default=20)
        ws.column_dimensions["A"].width = 8
        ws.column_dimensions["B"].width = max(28, max_kw_len + 6)
        ws.column_dimensions["C"].width = 20

        # Enable visual gridlines in Excel
        if ws.views.sheetView:
            ws.views.sheetView[0].showGridLines = True

        wb.save(xlsx_path)

        # Clean up legacy count_keywords.md if present, so output is strictly Excel
        old_md = output_dir / "count_keywords.md"
        if old_md.exists():
            try:
                os.remove(old_md)
            except OSError:
                pass

        print(f"[KEYWORDS] Exported Excel keyword report -> {xlsx_path}")
        return xlsx_path

    @classmethod
    def export_report(cls, report: KeywordReport, output_dir: Path) -> Path:
        """Standard export entrypoint returning the Excel spreadsheet."""
        return cls.export_excel(report, output_dir)

    @staticmethod
    def export_markdown(report: KeywordReport, output_dir: Path) -> Path:
        """
        Legacy Markdown exporter kept for backward compatibility with existing tests.
        """
        lines = [
            "# Báo Cáo Tần Suất Từ Khóa",
            "",
            f"**Tài liệu:** `{report.document_id}`  ",
            f"**Tổng số từ khóa tìm kiếm:** {report.total_keywords_searched}  ",
            f"**Khớp chính xác (100%):** {report.total_exact_matches}  ",
            f"**Tổng số lần phát hiện:** {report.total_matches}  ",
            "",
            "| Từ khóa / Cụm từ | Số lần xuất hiện |",
            "| :--- | :---: |",
        ]
        for kw in report.keywords:
            lines.append(f"| {kw.keyword} | {kw.count} |")

        lines.append("")

        out_path = Path(output_dir) / "count_keywords.md"
        with open(out_path, "w", encoding="utf-8") as f:
            f.write("\n".join(lines))

        return out_path

